# -*- coding: utf-8 -*-
"""
Robot de Quórum: lee el Excel de actas, resoluciones y normativa, lo revisa y
lo publica CIFRADO, un archivo por grupo de acceso.

    python3 scripts/publicar.py                       (en GitHub Actions)
    python3 scripts/publicar.py --excel libro.xlsx --salida _site/datos
    python3 scripts/publicar.py --demo                demo/datos-demo.json

Variables de entorno (secretos del repositorio):
    EXCEL_URL    enlace de descarga del Excel en OneDrive.
    EXCEL_CLAVE  contraseña con que está protegido el Excel.
    CLAVES       una línea por grupo de acceso:   legal = frase de acceso larga

Tres reglas que no se negocian, porque son actas de directorio:

1. Nada en claro. Sin frase para un grupo, ese grupo no se publica y el robot
   falla: mejor lo de ayer que algo legible.
2. Nada en los registros. El repositorio es público y los registros de GitHub
   Actions también: lo que se imprime aquí son conteos, nombres de hoja y
   números de fila, nunca el contenido de una celda. El detalle de cada aviso
   viaja dentro del archivo cifrado y se lee en «Revisión de la matriz».
3. Nada a medias. Si falta una hoja o una columna obligatoria, no se publica:
   queda en línea lo último que estuvo bien.
"""
import argparse
import base64
import datetime as dt
import io
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from matriz import (HOJAS, OBLIGATORIAS, TIPOS_CAMBIO, GRUPO_POR_OMISION,  # noqa: E402
                    plano, slug)

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ITER = 600000        # PBKDF2-SHA256: lo que recomienda OWASP desde 2023
MIN_FRASE = 20


class Alto(Exception):
    """Algo por lo que no se debe publicar. El mensaje va al registro: sin datos."""


# ----------------------------------------------------------------- lectura
def abrir_libro(contenido, clave, permitir_sin_clave):
    import openpyxl
    if contenido[:4] == b"\xd0\xcf\x11\xe0":           # contenedor OLE = Excel cifrado
        if not clave:
            raise Alto("El Excel está protegido con contraseña y falta el secreto EXCEL_CLAVE.")
        import msoffcrypto
        of = msoffcrypto.OfficeFile(io.BytesIO(contenido))
        salida = io.BytesIO()
        try:
            of.load_key(password=clave)
            of.decrypt(salida)
        except Exception:
            raise Alto("La contraseña de EXCEL_CLAVE no abre el Excel. ¿Se cambió la contraseña del libro?")
        contenido = salida.getvalue()
    elif contenido[:2] == b"PK":
        if not permitir_sin_clave:
            raise Alto("El Excel NO está protegido con contraseña. Como el robot lo baja con un enlace, "
                       "ese enlace daría las actas en claro a quien lo tenga. Protégelo en Excel: "
                       "Archivo → Información → Proteger libro → Cifrar con contraseña, y guarda la "
                       "contraseña en el secreto EXCEL_CLAVE.")
    else:
        raise Alto("Lo descargado no es un Excel. Revisa que EXCEL_URL apunte al archivo .xlsx.")
    return openpyxl.load_workbook(io.BytesIO(contenido), data_only=True)


def descargar(url):
    import requests
    if "download=1" not in url:
        url += ("&" if "?" in url else "?") + "download=1"
    r = requests.get(url, timeout=120, allow_redirects=True)
    if r.status_code != 200:
        raise Alto(f"No se pudo bajar el Excel (HTTP {r.status_code}). Revisa el enlace de EXCEL_URL.")
    return r.content


def mapa_columnas(ws, nombre):
    """Encabezado → clave. Reconoce el encabezado sin tildes, en otro orden y con
    algo añadido al final («Fecha de la sesión (dd/mm)»)."""
    encs = [plano(c.value) for c in ws[1]]
    mapa, sin = {}, []
    for clave, enc, *_ in HOJAS[nombre]:
        e = plano(enc)
        j = next((j for j, h in enumerate(encs) if h == e), None)
        if j is None:
            j = next((j for j, h in enumerate(encs) if h and h.startswith(e) and j not in mapa.values()), None)
        if j is None:
            sin.append(enc)
        else:
            mapa[clave] = j
    falta = [enc for clave, enc, *_ in HOJAS[nombre] if clave in OBLIGATORIAS[nombre] and clave not in mapa]
    if falta:
        raise Alto(f"En la hoja «{nombre}» falta la columna «{falta[0]}». No se publica hasta que vuelva.")
    return mapa


def filas(wb, nombre):
    if nombre not in wb.sheetnames:
        raise Alto(f"Falta la hoja «{nombre}» en el Excel. No se publica hasta que vuelva.")
    ws = wb[nombre]
    mapa = mapa_columnas(ws, nombre)
    for i, fila in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        d = {k: (fila[j] if j < len(fila) else None) for k, j in mapa.items()}
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in d.values()):
            continue
        yield i, d


def texto(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%d/%m/%Y")
    s = re.sub(r"\s+", " ", str(v)).strip()
    return s or None


def fecha(v):
    """ISO aaaa-mm-dd, o ValueError. Vacío → None."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    s = str(v).strip()
    for f in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return dt.datetime.strptime(s, f).date().isoformat()
        except ValueError:
            pass
    raise ValueError(s)


def fh(iso):
    """aaaa-mm-dd → dd/mm/aaaa, como se escribe en el Excel (para los avisos)."""
    return f"{iso[8:10]}/{iso[5:7]}/{iso[:4]}" if iso else ""


def punto(v):
    """«3», 3, 3.0, «Punto 3», «3.» → «3». Lo que no es número se deja como texto."""
    s = texto(v)
    if not s:
        return None
    s = re.sub(r"^(punto|pto\.?)\s*", "", s, flags=re.I).strip().rstrip(".")
    try:
        f = float(s.replace(",", "."))
        if f.is_integer():
            return str(int(f))
    except ValueError:
        pass
    return s


def enlace(v):
    """Solo http(s). Un «javascript:» en una celda no puede llegar a un href."""
    s = texto(v)
    if not s:
        return None, False
    return (s, False) if re.match(r"^https?://", s, re.I) else (None, True)


# ------------------------------------------------------------- el modelo
class Matriz:
    def __init__(self):
        self.fondos, self.cuerpos, self.actas, self.res, self.normas, self.cambios = {}, {}, {}, {}, {}, []
        self.avisos = []

    def aviso(self, hoja, fila, generico, detalle=None, grupo=None):
        """`generico` va al registro público; `detalle` solo al archivo cifrado."""
        self.avisos.append({"hoja": hoja, "fila": fila, "texto": detalle or generico,
                            "generico": generico, "grupo": grupo})

    def grupo_de(self, sigla):
        return self.cuerpos[sigla]["grupo"]

    def ref_res(self, sigla, f, p):
        return f"{slug(sigla)}_{f}_p{slug(p) or 'x'}"

    def leer(self, wb):
        self._fondos(wb)
        self._cuerpos(wb)
        self._actas(wb)
        self._resoluciones(wb)
        self._normas(wb)
        self._cambios(wb)
        return self

    def _sigla(self, hoja, i, v, rotulo="cuerpo"):
        s = texto(v)
        if not s:
            return None
        s = s.upper()
        if s not in self.cuerpos:
            self.aviso(hoja, i, f"el {rotulo} no está en la hoja Cuerpos; el renglón no se publica",
                       f"«{s}» no está en la hoja Cuerpos. Añádelo allí o corrige la sigla: mientras tanto el renglón no se publica.")
            return False
        return s

    def _fecha(self, hoja, i, v, rotulo, obligatoria=True, grupo=None):
        try:
            f = fecha(v)
        except ValueError:
            self.aviso(hoja, i, f"{rotulo}: no es una fecha", f"{rotulo}: «{texto(v)}» no es una fecha (dd/mm/aaaa).", grupo)
            return False
        if f is None and obligatoria:
            self.aviso(hoja, i, f"falta {rotulo}", None, grupo)
            return False
        return f

    def _fondos(self, wb):
        """La hoja es opcional: un Excel sin ella sigue publicándose, con el fondo
        tal como esté escrito en Cuerpos y Normativa."""
        if "Fondos" not in wb.sheetnames:
            return
        for i, d in filas(wb, "Fondos"):
            s = texto(d.get("sigla"))
            if not s:
                self.aviso("Fondos", i, "falta la sigla; el renglón no se publica")
                continue
            s = s.upper()
            if s in self.fondos:
                self.aviso("Fondos", i, "sigla repetida; vale el primer renglón", f"El fondo {s} está repetido; vale el primer renglón.")
                continue
            self.fondos[s] = {"sigla": s, "nombre": texto(d.get("nombre")) or s, "tipo": texto(d.get("tipo"))}

    def _fondo(self, hoja, i, v, grupo=None):
        """La sigla del fondo, como está en la hoja Fondos. Una que no está se
        avisa pero se respeta: es solo una etiqueta, no decide quién ve qué."""
        s = texto(v)
        if not s:
            return None
        s = s.upper()
        if self.fondos and s not in self.fondos:
            self.aviso(hoja, i, "el fondo no está en la hoja Fondos",
                       f"«{s}» no está en la hoja Fondos. Añádelo allí para que se vea con su nombre completo.", grupo)
        return s

    def _cuerpos(self, wb):
        for i, d in filas(wb, "Cuerpos"):
            s = texto(d.get("sigla"))
            if not s:
                self.aviso("Cuerpos", i, "falta la sigla; el renglón no se publica")
                continue
            s = s.upper()
            if s in self.cuerpos:
                self.aviso("Cuerpos", i, "sigla repetida; vale el primer renglón", f"La sigla {s} está repetida; vale el primer renglón.")
                continue
            self.cuerpos[s] = {"sigla": s, "nombre": texto(d.get("nombre")) or s,
                               "entidad": self._fondo("Cuerpos", i, d.get("entidad")),
                               "grupo": slug(texto(d.get("grupo")) or "") or GRUPO_POR_OMISION}

    def _actas(self, wb):
        for i, d in filas(wb, "Actas"):
            s = self._sigla("Actas", i, d.get("cuerpo"))
            if not s:
                if s is None:
                    self.aviso("Actas", i, "falta el cuerpo colegiado; el renglón no se publica")
                continue
            g = self.grupo_de(s)
            f = self._fecha("Actas", i, d.get("fecha"), "la fecha de la sesión", grupo=g)
            if not f:
                continue
            aid = f"{slug(s)}_{f}"
            if aid in self.actas:
                self.aviso("Actas", i, "acta repetida (mismo cuerpo y fecha); vale el primer renglón",
                           f"Ya hay un acta de {s} del {fh(f)}: vale el primer renglón. Si fueron dos sesiones el mismo día, avisa: Quórum las identifica por cuerpo y fecha.", g)
                continue
            a = {"id": aid, "cuerpo": s, "fecha": f, "tipo": texto(d.get("tipo")),
                 "numero": texto(d.get("numero")), "obs": texto(d.get("obs")), "fila": i}
            for k, rot in (("provisional", "acta provisional"), ("definitiva", "acta definitiva")):
                u, malo = enlace(d.get(k))
                if malo:
                    self.aviso("Actas", i, f"el enlace del {rot} no empieza por https://; se ignora", None, g)
                a[k] = u
            ap = self._fecha("Actas", i, d.get("aprobacion"), "la fecha de aprobación del acta definitiva", False, g)
            a["aprobacion"] = ap or None
            if ap and ap < f:
                self.aviso("Actas", i, "el acta definitiva se aprobó antes de la sesión",
                           f"El acta de {s} del {fh(f)} dice aprobarse el {fh(ap)}, antes de la propia sesión.", g)
            self.actas[aid] = a

    def _resoluciones(self, wb):
        for i, d in filas(wb, "Resoluciones"):
            s = self._sigla("Resoluciones", i, d.get("cuerpo"))
            if not s:
                if s is None:
                    self.aviso("Resoluciones", i, "falta el cuerpo colegiado; el renglón no se publica")
                continue
            g = self.grupo_de(s)
            f = self._fecha("Resoluciones", i, d.get("fecha"), "la fecha de la sesión", grupo=g)
            if not f:
                continue
            p, t = punto(d.get("punto")), texto(d.get("texto"))
            if not p or not t:
                self.aviso("Resoluciones", i, "falta el punto o lo resuelto; el renglón no se publica", None, g)
                continue
            aid = f"{slug(s)}_{f}"
            if aid not in self.actas:
                self.actas[aid] = {"id": aid, "cuerpo": s, "fecha": f, "stub": True}
                self.aviso("Resoluciones", i, "la sesión no está en la hoja Actas",
                           f"La resolución de {s} del {fh(f)} (punto {p}) no tiene su acta en la hoja Actas. Se muestra igual, con el acta incompleta.", g)
            rid = self.ref_res(s, f, p)
            if rid in self.res:
                self.aviso("Resoluciones", i, "resolución repetida (mismo cuerpo, fecha y punto); vale el primer renglón",
                           f"{s} · {fh(f)} · punto {p} está dos veces; vale el primer renglón.", g)
                continue
            u, malo = enlace(d.get("enlace"))
            if malo:
                self.aviso("Resoluciones", i, "el enlace no empieza por https://; se ignora", None, g)
            self.res[rid] = {"id": rid, "acta": aid, "cuerpo": s, "fecha": f, "punto": p, "texto": t,
                             "materia": texto(d.get("materia")), "enlace": u, "obs": texto(d.get("obs")), "fila": i}

    def _normas(self, wb):
        self.por_nombre = {}
        for i, d in filas(wb, "Normativa"):
            n = texto(d.get("nombre"))
            if not n:
                self.aviso("Normativa", i, "falta el nombre de la norma; el renglón no se publica")
                continue
            s = self._sigla("Normativa", i, d.get("organo"), "órgano que la aprueba")
            if not s:
                if s is None:
                    self.aviso("Normativa", i, "falta el órgano que la aprueba; el renglón no se publica",
                               f"«{n}»: falta el órgano que la aprueba. Sin él no se sabe quién puede verla, así que no se publica.")
                continue
            g = slug(texto(d.get("grupo")) or "") or self.grupo_de(s)
            f = self._fecha("Normativa", i, d.get("fecha"), "la fecha de aprobación", grupo=g)
            if not f:
                continue
            nid = slug(n)
            if nid in self.normas:
                self.aviso("Normativa", i, "norma repetida; vale el primer renglón", f"«{n}» está dos veces; vale el primer renglón.", g)
                continue
            p = punto(d.get("punto"))
            res = self.ref_res(s, f, p) if p else None
            if res and res not in self.res:
                self.aviso("Normativa", i, "no encuentro la resolución que la aprobó",
                           f"«{n}»: no encuentro en Resoluciones el punto {p} de {s} del {fh(f)}, que la aprobó.", g)
                res = None
            u, malo = enlace(d.get("enlace"))
            if malo:
                self.aviso("Normativa", i, "el enlace no empieza por https://; se ignora", None, g)
            aid = f"{slug(s)}_{f}"
            self.normas[nid] = {"id": nid, "nombre": n, "tipo": texto(d.get("tipo")),
                                "entidad": self._fondo("Normativa", i, d.get("entidad"), g) or self.cuerpos[s]["entidad"],
                                "organo": s, "fecha": f, "punto": p, "res": res,
                                "acta": aid if aid in self.actas else None, "enlace": u,
                                "obs": texto(d.get("obs")), "grupo": g, "fila": i}
            self.por_nombre[plano(n)] = nid

    def _cambios(self, wb):
        tipos = {plano(t): t for t in TIPOS_CAMBIO}
        for i, d in filas(wb, "Cambios"):
            tipo = tipos.get(plano(d.get("tipo")))
            if not tipo:
                self.aviso("Cambios", i, "tipo de cambio vacío o fuera de la lista; el renglón no se publica")
                continue
            nombre = texto(d.get("norma"))
            hay_res = any(texto(d.get(k)) for k in ("rcuerpo", "rfecha", "rpunto"))
            if bool(nombre) == hay_res:
                self.aviso("Cambios", i, "hay que llenar la norma afectada O la resolución afectada (una de las dos)")
                continue
            if nombre:
                nid = self.por_nombre.get(plano(nombre))
                if not nid:
                    self.aviso("Cambios", i, "la norma afectada no está en la hoja Normativa; el renglón no se publica",
                               f"«{nombre}» no está en la hoja Normativa (¿el nombre está escrito distinto?).")
                    continue
                n = self.normas[nid]
                a = {"t": "norma", "id": nid, "nombre": n["nombre"], "fecha": n["fecha"]}
                grupos = {n["grupo"]}
            else:
                s = self._sigla("Cambios", i, d.get("rcuerpo"))
                if not s:
                    continue
                f = self._fecha("Cambios", i, d.get("rfecha"), "la fecha de la resolución afectada", grupo=self.grupo_de(s))
                p = punto(d.get("rpunto"))
                if not f or not p:
                    if f and not p:
                        self.aviso("Cambios", i, "falta el punto de la resolución afectada", None, self.grupo_de(s))
                    continue
                rid = self.ref_res(s, f, p)
                if rid not in self.res:
                    self.aviso("Cambios", i, "la resolución afectada no está en la hoja Resoluciones; el renglón no se publica",
                               f"No encuentro en Resoluciones el punto {p} de {s} del {fh(f)}.", self.grupo_de(s))
                    continue
                a = {"t": "res", "id": rid, "cuerpo": s, "fecha": f, "punto": p}
                grupos = {self.grupo_de(s)}
            g0 = sorted(grupos)[0]

            por = None
            if any(texto(d.get(k)) for k in ("pcuerpo", "pfecha", "ppunto")):
                s = self._sigla("Cambios", i, d.get("pcuerpo"), "cuerpo de «Resuelto por»")
                if s is False:
                    continue
                f = self._fecha("Cambios", i, d.get("pfecha"), "la fecha de «Resuelto por»", grupo=g0)
                if f is False:
                    continue
                p = punto(d.get("ppunto"))
                if not s or not f or not p:
                    self.aviso("Cambios", i, "«Resuelto por» está incompleto: van cuerpo, fecha y punto", None, g0)
                    continue
                rid = self.ref_res(s, f, p)
                por = {"id": rid if rid in self.res else None, "cuerpo": s, "fecha": f, "punto": p}
                if not por["id"]:
                    self.aviso("Cambios", i, "no encuentro la resolución de «Resuelto por»",
                               f"No encuentro en Resoluciones el punto {p} de {s} del {fh(f)}, que hizo el cambio. Se muestra igual, sin enlace.", g0)
                grupos.add(self.grupo_de(s))
            rige = self._fecha("Cambios", i, d.get("rige"), "«Rige desde»", False, g0)
            if rige is False:
                continue
            efecto = rige or (por and por["fecha"])
            if not efecto:
                self.aviso("Cambios", i, "sin «Resuelto por» ni «Rige desde» no hay fecha; el renglón no se publica", None, g0)
                continue
            if efecto < a["fecha"]:
                self.aviso("Cambios", i, "el cambio es anterior a lo que cambia",
                           f"El cambio rige el {fh(efecto)}, antes de la fecha de lo que cambia ({fh(a['fecha'])}).", g0)
            u, malo = enlace(d.get("enlace"))
            if malo:
                self.aviso("Cambios", i, "el enlace no empieza por https://; se ignora", None, g0)
            self.cambios.append({"id": f"c{i}", "tipo": tipo, "fecha": efecto, "rige": rige, "a": a, "por": por,
                                 "articulos": texto(d.get("articulos")), "enlace": u,
                                 "detalle": texto(d.get("detalle")), "grupos": sorted(grupos), "fila": i})

    # --------------------------------------------------------- reparto
    def grupos(self):
        gs = {c["grupo"] for c in self.cuerpos.values()} | {n["grupo"] for n in self.normas.values()}
        return sorted(gs)

    def payload(self, grupo, generado):
        def de(sigla):
            return self.grupo_de(sigla) == grupo
        limpio = lambda d: {k: v for k, v in d.items() if k not in ("fila", "grupos") and v not in (None, "", [])}
        normas = [limpio(n) for n in self.normas.values() if n["grupo"] == grupo]
        usados = {c["entidad"] for c in self.cuerpos.values()} | {n.get("entidad") for n in normas}
        return {
            "v": 1, "grupo": grupo, "generado": generado,
            "cuerpos": [{"sigla": c["sigla"], "nombre": c["nombre"], "entidad": c["entidad"]} for c in self.cuerpos.values()],
            "fondos": [limpio(f) for f in self.fondos.values() if f["sigla"] in usados],
            "actas": [limpio(a) for a in self.actas.values() if de(a["cuerpo"])],
            "resoluciones": [limpio(r) for r in self.res.values() if de(r["cuerpo"])],
            "normas": normas,
            "cambios": [limpio(c) for c in self.cambios if grupo in c["grupos"]],
            "avisos": [{"hoja": a["hoja"], "fila": a["fila"], "texto": a["texto"]}
                       for a in self.avisos if a["grupo"] in (None, grupo)],
        }


# ------------------------------------------------------------- cifrado
def cifrar(datos, frase):
    """AES-256-GCM con clave derivada de la frase (PBKDF2-SHA256). Lo descifra
    WebCrypto en el navegador, con los mismos parámetros."""
    from cryptography.hazmat.primitives import hashes
    from cryptography.hazmat.primitives.ciphers.aead import AESGCM
    from cryptography.hazmat.primitives.kdf.pbkdf2 import PBKDF2HMAC
    salt, iv = os.urandom(16), os.urandom(12)
    clave = PBKDF2HMAC(algorithm=hashes.SHA256(), length=32, salt=salt, iterations=ITER).derive(frase.encode("utf-8"))
    ct = AESGCM(clave).encrypt(iv, json.dumps(datos, ensure_ascii=False, separators=(",", ":")).encode("utf-8"), None)
    b = lambda x: base64.b64encode(x).decode("ascii")
    return {"q_enc": 1, "kdf": "PBKDF2-SHA256", "iter": ITER, "salt": b(salt), "iv": b(iv), "ct": b(ct)}


def leer_claves(texto_claves):
    claves = {}
    for linea in (texto_claves or "").splitlines():
        linea = linea.strip()
        if not linea or linea.startswith("#"):
            continue
        m = re.match(r"^([^=:]+)[=:](.+)$", linea)
        if not m:
            raise Alto("CLAVES tiene una línea que no es «grupo = frase». Revisa el secreto.")
        claves[slug(m.group(1))] = m.group(2).strip()
    return claves


def escribir(ruta, obj):
    os.makedirs(os.path.dirname(ruta), exist_ok=True)
    with open(ruta, "w", encoding="utf-8") as f:
        json.dump(obj, f, ensure_ascii=False, separators=(",", ":"))


def resumen(m, grupos):
    lineas = [f"Cuerpos: {len(m.cuerpos)} · fondos: {len(m.fondos)} · actas: {sum(1 for a in m.actas.values() if not a.get('stub'))} · "
              f"resoluciones: {len(m.res)} · normas: {len(m.normas)} · cambios: {len(m.cambios)}",
              f"Grupos publicados: {', '.join(grupos)}"]
    if m.avisos:
        lineas.append(f"Avisos: {len(m.avisos)} (el detalle está en Quórum → Revisión de la matriz)")
        lineas += [f"  · {a['hoja']}, fila {a['fila']}: {a['generico']}" for a in m.avisos[:40]]
    return "\n".join(lineas)


def main(argv=None):
    ap = argparse.ArgumentParser(description="Publica el Excel de Quórum, cifrado.")
    ap.add_argument("--excel", help="Excel local, en vez de bajarlo de EXCEL_URL")
    ap.add_argument("--salida", default=os.path.join(RAIZ, "_site", "datos"))
    ap.add_argument("--demo", action="store_true", help="publica demo/matriz_demo.xlsx SIN cifrar (solo datos ficticios)")
    ap.add_argument("--permitir-sin-clave", action="store_true", help="acepta un Excel sin contraseña (pruebas)")
    ap.add_argument("--si-falta-config", action="store_true",
                    help="si aún no hay secretos, publica la herramienta sin datos en vez de fallar")
    a = ap.parse_args(argv)
    ahora = dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat()

    if a.demo:
        ruta = a.excel or os.path.join(RAIZ, "demo", "matriz_demo.xlsx")
        with open(ruta, "rb") as f:
            wb = abrir_libro(f.read(), None, True)
        if "FICTICIOS" not in (wb.properties.title or ""):
            raise Alto("--demo solo publica el libro de ejemplo (el que crea crear_plantilla.py --demo). "
                       "Un Excel real nunca se publica sin cifrar.")
        m = Matriz().leer(wb)
        datos = m.payload(GRUPO_POR_OMISION, ahora)
        datos.update({"demo": True, "grupo": "demo"})
        escribir(os.path.join(RAIZ, "demo", "datos-demo.json"), datos)
        print(resumen(m, ["demo (sin cifrar: datos ficticios)"]))
        return 0

    url = os.environ.get("EXCEL_URL", "").strip()
    if not a.excel and not url:
        if a.si_falta_config:
            escribir(os.path.join(a.salida, "indice.json"), {"generado": ahora, "grupos": [], "pendiente": "configuracion"})
            print("Aún no hay EXCEL_URL: se publica la herramienta sin datos (solo la demostración).")
            return 0
        raise Alto("Falta el secreto EXCEL_URL (o --excel para un archivo local).")

    claves = leer_claves(os.environ.get("CLAVES", ""))
    if a.excel:
        with open(a.excel, "rb") as f:
            contenido = f.read()
    else:
        contenido = descargar(url)
    wb = abrir_libro(contenido, os.environ.get("EXCEL_CLAVE", "").strip(), a.permitir_sin_clave)
    m = Matriz().leer(wb)
    grupos = m.grupos()

    faltan = [g for g in grupos if g not in claves]
    if faltan:
        raise Alto("No hay frase en el secreto CLAVES para: " + ", ".join(faltan) +
                   ". No se publica nada en claro. Añade una línea «grupo = frase» por cada uno.")
    cortas = [g for g in grupos if len(claves[g]) < MIN_FRASE]
    if cortas:
        raise Alto(f"La frase de {', '.join(cortas)} es corta (menos de {MIN_FRASE} caracteres). "
                   "Genera una en Quórum → «Generar una frase nueva».")

    indice = {"generado": ahora, "grupos": []}
    for g in grupos:
        archivo = f"g-{g}.json"
        escribir(os.path.join(a.salida, archivo), cifrar(m.payload(g, ahora), claves[g]))
        indice["grupos"].append({"id": g, "archivo": archivo})
    escribir(os.path.join(a.salida, "indice.json"), indice)

    txt = resumen(m, grupos)
    print(txt)
    if os.environ.get("GITHUB_STEP_SUMMARY"):
        with open(os.environ["GITHUB_STEP_SUMMARY"], "a", encoding="utf-8") as f:
            f.write("### Quórum publicado\n\n```\n" + txt + "\n```\n")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Alto as e:
        print("ERROR: " + str(e), file=sys.stderr)
        sys.exit(1)
