# -*- coding: utf-8 -*-
"""
Arma el Excel de Quórum: la plantilla vacía que se llena, y el de ejemplo.

    python3 scripts/crear_plantilla.py            plantilla/Quorum_matriz.xlsx
    python3 scripts/crear_plantilla.py --demo     demo/matriz_demo.xlsx (datos ficticios)

La forma de cada hoja sale de `matriz.py`, la misma que lee el robot. Los datos
de ejemplo son INVENTADOS: sirven para ver la herramienta funcionando antes de
llenar el Excel de verdad, y llevan la marca FICTICIOS en las propiedades del
libro, que es lo que `publicar.py --demo` exige para publicarlos sin cifrar.
"""
import os
import sys
import datetime as dt

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from matriz import HOJAS, LISTAS, CUERPOS_INICIALES  # noqa: E402

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
NAVY = "00249C"
MARCA_DEMO = "Quórum · DATOS FICTICIOS DE EJEMPLO"
FILAS = 2000  # hasta dónde llegan las listas desplegables

INSTRUCCIONES = [
    ("titulo", "Quórum · Matriz de actas, resoluciones y normativa"),
    ("", "Este libro es la fuente de Quórum. Lo que se escribe aquí es lo que se ve en la herramienta, cifrado, después de que el robot lo publica."),
    ("seccion", "Cómo se llena"),
    ("", "1. Cuerpos: un renglón por cada cuerpo colegiado. Ya vienen el Directorio del FIAS, la Junta Administrativa del FAP y el Directorio del FEIG; se pueden añadir más."),
    ("", "2. Programas: un renglón por cada programa o proyecto. Con ellos se llena la columna «Programa o proyecto» de Actas y Resoluciones, y así se puede consultar todo lo de un proyecto junto. Si algo es general (no es de ningún proyecto), esa columna se deja vacía."),
    ("", "3. Actas: un renglón por sesión. Cuerpo + fecha de la sesión identifican al acta. Si la sesión trata de uno o varios programas o proyectos, se escriben en «Programa o proyecto», separados con punto y coma. Mientras no tenga enlace al acta definitiva ni fecha de aprobación, se muestra como provisional."),
    ("", "4. Resoluciones: un renglón por cada punto en que se resolvió algo. Cuerpo + fecha + punto la identifican: «en la Junta Administrativa del 8 de agosto de 2026, en el tercer punto» se escribe JA-FAP · 08/08/2026 · 3. Por omisión es del programa o proyecto de su acta; la columna de programa solo se llena si trata de otro."),
    ("", "5. Normativa: un renglón por norma, con el órgano que la aprueba y la fecha (y si se sabe, el punto) en que se aprobó."),
    ("", "6. Cambios: un renglón cada vez que una resolución reforma, codifica, deroga, deja sin efecto o aclara una norma u otra resolución."),
    ("seccion", "La vigencia no se escribe: se deduce"),
    ("", "No hay columna «vigente». Quórum la calcula de la hoja Cambios, a la fecha que se pregunte: una norma con una Derogación es derogada desde esa fecha; con una Reforma o Codificación, sigue vigente con texto nuevo."),
    ("", "Cuando un acta nueva da de baja algo, solo se añade un renglón en Cambios. El renglón viejo no se toca: así nunca hay dos lugares que se contradigan."),
    ("", "Si el cambio rige desde otra fecha (p. ej. «a partir del 1 de enero»), va en «Rige desde». Hasta entonces, Quórum lo muestra como cambio próximo."),
    ("seccion", "Ejemplo de un renglón en Cambios"),
    ("", "Norma afectada: Manual Operativo del Ciclo de Proyectos del FAP · Tipo de cambio: Reforma · Resuelto por: JA-FAP · 08/08/2026 · 3 · Texto vigente: enlace al manual reformado."),
    ("seccion", "Confidencialidad: tres reglas"),
    ("", "1. Este libro se protege con contraseña, como un PDF: en Excel de escritorio, Archivo → Información → Proteger libro → Cifrar con contraseña. Es una contraseña que tú inventas y que se guarda en GitHub como secreto (EXCEL_CLAVE): con ella el robot abre el libro. Sin ella, el enlace de descarga solo entrega un archivo cerrado. Este libro se edita en Excel de escritorio (Excel en el navegador no abre libros con contraseña)."),
    ("", "2. Los enlaces a los PDF deben ser de la carpeta compartida solo con quien corresponde, nunca «cualquier persona con el vínculo»."),
    ("", "3. No se cambia el nombre de las hojas ni de los encabezados de la fila 1. Las columnas sí se pueden reordenar o ensanchar."),
    ("seccion", "Publicar"),
    ("", "El robot revisa este libro todos los días y publica lo que encuentre. Para publicar en el momento: en GitHub, pestaña Actions → «Publicar Quórum» → Run workflow."),
    ("", "Si algo no cuadra (una resolución sin su acta, un cambio que apunta a una norma que no existe), Quórum lo muestra en «Revisión de la matriz», con la hoja y el renglón."),
]


def hoja_listas(wb, cuerpos_rango, normas_rango):
    ws = wb.create_sheet("Listas")
    rangos = {"CUERPOS": cuerpos_rango, "NORMAS": normas_rango}
    for j, (nombre, valores) in enumerate(LISTAS.items(), start=1):
        col = get_column_letter(j)
        ws.cell(row=1, column=j, value=nombre).font = Font(bold=True)
        for i, v in enumerate(valores, start=2):
            ws.cell(row=i, column=j, value=v)
        rangos[nombre] = f"Listas!${col}$2:${col}${len(valores) + 1}"
        ws.column_dimensions[col].width = 26
    ws.sheet_state = "hidden"
    return rangos


def armar_hoja(wb, nombre, filas, rangos):
    cols = HOJAS[nombre]
    ws = wb.create_sheet(nombre)
    for j, (clave, enc, ayuda, ancho, tipo) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=j, value=enc)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        if ayuda:
            c.comment = Comment(ayuda, "Quórum", width=260, height=110)
        ws.column_dimensions[get_column_letter(j)].width = ancho
    ws.row_dimensions[1].height = 46
    ws.freeze_panes = "A2"

    for i, fila in enumerate(filas, start=2):
        for j, (clave, _e, _a, _w, tipo) in enumerate(cols, start=1):
            v = fila.get(clave)
            if v in (None, ""):
                continue
            if tipo == "fecha" and isinstance(v, str):
                v = dt.datetime.strptime(v, "%Y-%m-%d")
            ws.cell(row=i, column=j, value=v)

    for j, (clave, _e, _a, _w, tipo) in enumerate(cols, start=1):
        col = get_column_letter(j)
        rango = f"{col}2:{col}{FILAS}"
        for i in range(2, max(len(filas) + 2, 3)):
            celda = ws.cell(row=i, column=j)
            if tipo == "fecha":
                celda.number_format = "dd/mm/yyyy"
            elif tipo == "texto" and _w >= 40:
                celda.alignment = Alignment(wrap_text=True, vertical="top")
        if tipo == "fecha":
            dv = DataValidation(type="date", operator="greaterThan", formula1="DATE(1990,1,1)",
                                allow_blank=True, showErrorMessage=True,
                                errorTitle="Fecha", error="Escribe una fecha: dd/mm/aaaa.")
        elif tipo.startswith("lista:"):
            lista = tipo.split(":", 1)[1]
            abierta = lista in ("MATERIAS", "TIPOS_NORMA")
            dv = DataValidation(type="list", formula1=rangos[lista], allow_blank=True,  # sin «=»: así lo escribe Excel
                                showErrorMessage=not abierta,
                                errorTitle="Valor de la lista",
                                error="Elige un valor de la lista. Si es un cuerpo o una norma nueva, añádelo primero en su hoja.")
        else:
            continue
        ws.add_data_validation(dv)
        dv.add(rango)

    ultima = max(len(filas) + 1, 2)
    t = Table(displayName="t" + nombre, ref=f"A1:{get_column_letter(len(cols))}{ultima}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(t)
    return ws


def instrucciones(wb, demo):
    ws = wb.active
    ws.title = "Instrucciones"
    ws.column_dimensions["A"].width = 120
    fila = 1
    if demo:
        c = ws.cell(row=fila, column=1, value="DATOS FICTICIOS DE EJEMPLO — no corresponden a ninguna sesión real.")
        c.font = Font(bold=True, color="C0392B", size=13)
        fila += 2
    for tipo, txt in INSTRUCCIONES:
        if tipo == "seccion":
            fila += 1   # un renglón en blanco antes de cada sección
        c = ws.cell(row=fila, column=1, value=txt)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if tipo == "titulo":
            c.font = Font(bold=True, size=16, color=NAVY)
        elif tipo == "seccion":
            c.font = Font(bold=True, size=12, color=NAVY)
        fila += 1
    ws.sheet_view.showGridLines = False


def libro(demo=False):
    datos = demo_datos() if demo else {"Cuerpos": [
        {"sigla": s, "nombre": n, "entidad": e, "grupo": g} for s, n, e, g in CUERPOS_INICIALES]}
    wb = Workbook()
    instrucciones(wb, demo)
    rangos = hoja_listas(wb, f"Cuerpos!$A$2:$A${FILAS}", f"Normativa!$A$2:$A${FILAS}")
    for nombre in HOJAS:
        armar_hoja(wb, nombre, datos.get(nombre, []), rangos)
    # Las listas al final, para que la AC no tropiece con ellas.
    wb.move_sheet("Listas", offset=len(wb.sheetnames))
    wb.properties.title = MARCA_DEMO if demo else "Quórum · Matriz"
    wb.properties.creator = "Quórum"
    return wb


# --------------------------------------------------------------------------
# Datos de ejemplo. TODO ES INVENTADO: nombres de normas plausibles para que la
# demostración se entienda, pero ninguna fecha, punto ni texto es real.
# --------------------------------------------------------------------------
EJ = "https://ejemplo.invalid/quorum/"

SESIONES = {
    "DIR-FIAS": [("2017-05-18", "Ordinaria"), ("2017-11-23", "Ordinaria"), ("2018-08-16", "Ordinaria"),
                 ("2019-03-14", "Ordinaria"), ("2019-07-11", "Ordinaria"), ("2019-12-12", "Ordinaria"),
                 ("2021-03-18", "Ordinaria"), ("2022-06-09", "Ordinaria"), ("2023-04-20", "Extraordinaria"),
                 ("2024-01-25", "Ordinaria"), ("2024-10-17", "Ordinaria"), ("2025-03-20", "Ordinaria"),
                 ("2025-11-20", "Ordinaria"), ("2026-09-17", "Ordinaria")],
    "JA-FAP": [("2017-03-09", "Ordinaria"), ("2018-02-20", "Ordinaria"), ("2020-09-10", "Ordinaria"),
               ("2021-08-12", "Ordinaria"), ("2024-06-13", "Ordinaria"), ("2024-09-12", "Extraordinaria"),
               ("2026-08-08", "Ordinaria"), ("2026-10-01", "Ordinaria")],
    "DIR-FEIG": [("2020-02-13", "Ordinaria"), ("2025-05-15", "Ordinaria"), ("2026-05-14", "Ordinaria")],
}

RESOLUCIONES = [
    ("DIR-FIAS", "2017-05-18", "3", "Conocer y aprobar el informe de gestión del primer trimestre de 2017.", "Gobierno corporativo"),
    ("DIR-FIAS", "2017-05-18", "4", "Aprobar el Reglamento de Adquisiciones del FIAS.", "Normativa"),
    ("DIR-FIAS", "2017-11-23", "2", "Aprobar la Política de Inversiones del FIAS.", "Inversiones"),
    ("DIR-FIAS", "2017-11-23", "3", "Aprobar el presupuesto institucional 2018.", "Presupuesto y finanzas"),
    ("DIR-FIAS", "2018-08-16", "3", "Designar al auditor externo para el ejercicio 2018.", "Designaciones"),
    ("DIR-FIAS", "2018-08-16", "5", "Aprobar el Código de Ética y Conducta del FIAS.", "Normativa"),
    ("DIR-FIAS", "2019-03-14", "3", "Reformar los artículos 12 y 15 del Reglamento de Adquisiciones para actualizar los montos de contratación directa.", "Normativa"),
    ("DIR-FIAS", "2019-07-11", "6", "Aprobar el Reglamento Interno de Talento Humano.", "Talento humano"),
    ("DIR-FIAS", "2019-12-12", "2", "Aprobar el presupuesto institucional 2020.", "Presupuesto y finanzas"),
    ("DIR-FIAS", "2019-12-12", "4", "Aprobar el Procedimiento de Viáticos y Movilización.", "Normativa"),
    ("DIR-FIAS", "2021-03-18", "5", "Aprobar los Lineamientos de Teletrabajo.", "Talento humano"),
    ("DIR-FIAS", "2022-06-09", "2", "Aprobar la codificación del Reglamento de Adquisiciones del FIAS, que integra sus reformas.", "Normativa"),
    ("DIR-FIAS", "2023-04-20", "4", "Aprobar la Política de Inversiones y Gestión de Portafolio y derogar la Política de Inversiones de 2017.", "Inversiones"),
    ("DIR-FIAS", "2024-01-25", "3", "Reformar el capítulo V del Reglamento Interno de Talento Humano para integrar el régimen de viáticos.", "Talento humano"),
    ("DIR-FIAS", "2024-01-25", "4", "Derogar el Procedimiento de Viáticos y Movilización, cuyo contenido pasa al Reglamento Interno de Talento Humano.", "Normativa"),
    ("DIR-FIAS", "2024-10-17", "2", "Aclarar el alcance del artículo 8 de la Política de Inversiones y Gestión de Portafolio sobre los límites por emisor.", "Inversiones"),
    ("DIR-FIAS", "2025-03-20", "3", "Designar a la Gerencia Financiera como representante del FIAS ante el comité de inversiones.", "Designaciones"),
    ("DIR-FIAS", "2025-11-20", "2", "Reformar la designación del 20 de marzo de 2025: la representación ante el comité de inversiones la ejerce la Dirección Ejecutiva.", "Designaciones"),
    ("DIR-FIAS", "2026-09-17", "2", "Derogar los Lineamientos de Teletrabajo a partir del 1 de enero de 2027.", "Talento humano"),
    ("DIR-FIAS", "2026-09-17", "3", "Aprobar el presupuesto institucional 2027.", "Presupuesto y finanzas"),
    ("JA-FAP", "2017-03-09", "3", "Aprobar el Reglamento de Funcionamiento de la Junta Administrativa del FAP.", "Gobierno corporativo"),
    ("JA-FAP", "2018-02-20", "5", "Aprobar el Manual Operativo del Ciclo de Proyectos del FAP.", "Normativa"),
    ("JA-FAP", "2020-09-10", "2", "Aprobar la codificación del Reglamento de Funcionamiento de la Junta Administrativa.", "Gobierno corporativo"),
    ("JA-FAP", "2021-08-12", "3", "Reformar la sección de seguimiento técnico del Manual Operativo del Ciclo de Proyectos.", "Normativa"),
    ("JA-FAP", "2024-06-13", "4", "Aprobar las bases de la convocatoria de pequeñas donaciones 2024, con cierre el 30 de septiembre.", "Proyectos y convocatorias"),
    ("JA-FAP", "2024-09-12", "2", "Dejar sin efecto la convocatoria de pequeñas donaciones 2024 por falta de disponibilidad presupuestaria.", "Proyectos y convocatorias"),
    ("JA-FAP", "2026-08-08", "3", "Aprobar la reforma integral del Manual Operativo del Ciclo de Proyectos del FAP.", "Normativa"),
    ("JA-FAP", "2026-08-08", "4", "Aprobar el plan operativo anual 2027 de las áreas protegidas.", "Presupuesto y finanzas"),
    ("JA-FAP", "2026-10-01", "2", "Conocer el informe de avance de la campaña de renovaciones 2027.", "Administrativo"),
    ("DIR-FEIG", "2020-02-13", "3", "Aprobar el Manual de Operaciones del FEIG.", "Normativa"),
    ("DIR-FEIG", "2025-05-15", "2", "Reformar el Manual de Operaciones del FEIG para incorporar los desembolsos por hitos, desde el 1 de julio de 2025.", "Normativa"),
    ("DIR-FEIG", "2026-05-14", "2", "Aprobar el informe financiero del primer trimestre de 2026.", "Presupuesto y finanzas"),
    ("DIR-FEIG", "2026-05-14", "3", "Aprobar el convenio de cooperación con una universidad para el monitoreo de especies.", "Convenios y alianzas"),
    # A propósito: una resolución cuya sesión no está en la hoja Actas, para que
    # la demostración muestre cómo avisa «Revisión de la matriz».
    ("DIR-FEIG", "2022-11-10", "2", "Ratificar la integración del comité técnico del FEIG.", "Gobierno corporativo"),
]

PROGRAMAS = [
    ("Programa de Pequeñas Donaciones", "Programa", "FAP"),
    ("Programa de Gestión de Áreas Protegidas", "Programa", "FAP"),
    ("Proyecto de Monitoreo de Especies", "Proyecto", "FEIG"),
]
PD, GAP, MON = (p[0] for p in PROGRAMAS)
# Programa o proyecto de cada sesión, y de las resoluciones que difieren de su acta.
PROG_ACTA = {("JA-FAP", "2024-06-13"): PD, ("JA-FAP", "2024-09-12"): PD,
             ("JA-FAP", "2026-08-08"): f"{PD}; {GAP}"}
PROG_RES = {("JA-FAP", "2026-08-08", "4"): GAP, ("DIR-FEIG", "2026-05-14", "3"): MON}

NORMAS = [
    ("Reglamento de Adquisiciones del FIAS", "Reglamento", "FIAS", "DIR-FIAS", "2017-05-18", "4"),
    ("Política de Inversiones del FIAS", "Política", "FIAS", "DIR-FIAS", "2017-11-23", "2"),
    ("Código de Ética y Conducta del FIAS", "Código", "FIAS", "DIR-FIAS", "2018-08-16", "5"),
    ("Reglamento Interno de Talento Humano", "Reglamento", "FIAS", "DIR-FIAS", "2019-07-11", "6"),
    ("Procedimiento de Viáticos y Movilización", "Procedimiento", "FIAS", "DIR-FIAS", "2019-12-12", "4"),
    ("Lineamientos de Teletrabajo", "Lineamiento", "FIAS", "DIR-FIAS", "2021-03-18", "5"),
    ("Política de Inversiones y Gestión de Portafolio", "Política", "FIAS", "DIR-FIAS", "2023-04-20", "4"),
    ("Reglamento de Funcionamiento de la Junta Administrativa del FAP", "Reglamento", "FAP", "JA-FAP", "2017-03-09", "3"),
    ("Manual Operativo del Ciclo de Proyectos del FAP", "Manual", "FAP", "JA-FAP", "2018-02-20", "5"),
    ("Manual de Operaciones del FEIG", "Manual", "FEIG", "DIR-FEIG", "2020-02-13", "3"),
]

CAMBIOS = [
    # (norma | (cuerpo, fecha, punto)), tipo, (por), rige, artículos, con texto nuevo
    ("Reglamento de Adquisiciones del FIAS", "Reforma parcial", ("DIR-FIAS", "2019-03-14", "3"), None, "arts. 12 y 15", True),
    ("Reglamento de Adquisiciones del FIAS", "Codificación", ("DIR-FIAS", "2022-06-09", "2"), None, None, True),
    ("Política de Inversiones del FIAS", "Derogación", ("DIR-FIAS", "2023-04-20", "4"), None, None, False),
    ("Política de Inversiones y Gestión de Portafolio", "Aclaración", ("DIR-FIAS", "2024-10-17", "2"), None, "art. 8", False),
    ("Reglamento Interno de Talento Humano", "Reforma parcial", ("DIR-FIAS", "2024-01-25", "3"), None, "capítulo V", True),
    ("Procedimiento de Viáticos y Movilización", "Derogación", ("DIR-FIAS", "2024-01-25", "4"), None, None, False),
    ("Lineamientos de Teletrabajo", "Derogación", ("DIR-FIAS", "2026-09-17", "2"), "2027-01-01", None, False),
    ("Reglamento de Funcionamiento de la Junta Administrativa del FAP", "Codificación", ("JA-FAP", "2020-09-10", "2"), None, None, True),
    ("Manual Operativo del Ciclo de Proyectos del FAP", "Reforma parcial", ("JA-FAP", "2021-08-12", "3"), None, "arts. 22 a 25", True),
    ("Manual Operativo del Ciclo de Proyectos del FAP", "Reforma", ("JA-FAP", "2026-08-08", "3"), None, None, True),
    ("Manual de Operaciones del FEIG", "Reforma parcial", ("DIR-FEIG", "2025-05-15", "2"), "2025-07-01", "sección de desembolsos", True),
    (("DIR-FIAS", "2025-03-20", "3"), "Reforma", ("DIR-FIAS", "2025-11-20", "2"), None, None, False),
    (("JA-FAP", "2024-06-13", "4"), "Deja sin efecto", ("JA-FAP", "2024-09-12", "2"), None, None, False),
]


def demo_datos():
    cuerpos = [{"sigla": s, "nombre": n, "entidad": e, "grupo": g} for s, n, e, g in CUERPOS_INICIALES]
    actas, resol = [], []
    for cuerpo, ses in SESIONES.items():
        for k, (fecha, tipo) in enumerate(ses):
            sig = ses[k + 1][0] if k + 1 < len(ses) else None
            a = {"cuerpo": cuerpo, "fecha": fecha, "tipo": tipo, "programa": PROG_ACTA.get((cuerpo, fecha)),
                 "provisional": f"{EJ}{cuerpo}/{fecha}-provisional.pdf"}
            if sig:   # la definitiva se aprueba en la sesión siguiente
                a["definitiva"] = f"{EJ}{cuerpo}/{fecha}-definitiva.pdf"
                a["aprobacion"] = sig
            actas.append(a)
            if k > 0:
                ant = dt.date.fromisoformat(ses[k - 1][0]).strftime("%d/%m/%Y")
                resol.append({"cuerpo": cuerpo, "fecha": fecha, "punto": "1",
                              "texto": f"Aprobar el acta de la sesión del {ant}.", "materia": "Gobierno corporativo"})
    for cuerpo, fecha, punto, texto, materia in RESOLUCIONES:
        resol.append({"cuerpo": cuerpo, "fecha": fecha, "punto": punto, "texto": texto, "materia": materia,
                      "programa": PROG_RES.get((cuerpo, fecha, punto))})
    normas = [{"nombre": n, "tipo": t, "entidad": e, "organo": o, "fecha": f, "punto": p,
               "enlace": f"{EJ}normativa/{i + 1}-original.pdf"} for i, (n, t, e, o, f, p) in enumerate(NORMAS)]
    cambios = []
    for i, (a, tipo, por, rige, arts, texto) in enumerate(CAMBIOS):
        c = {"tipo": tipo, "pcuerpo": por[0], "pfecha": por[1], "ppunto": por[2],
             "rige": rige, "articulos": arts}
        if isinstance(a, tuple):
            c.update({"rcuerpo": a[0], "rfecha": a[1], "rpunto": a[2]})
        else:
            c["norma"] = a
        if texto:
            c["enlace"] = f"{EJ}normativa/cambio-{i + 1}.pdf"
        cambios.append(c)
    progs = [{"nombre": n, "tipo": t, "entidad": e} for n, t, e in PROGRAMAS]
    return {"Cuerpos": cuerpos, "Programas": progs, "Actas": actas, "Resoluciones": resol, "Normativa": normas, "Cambios": cambios}


if __name__ == "__main__":
    demo = "--demo" in sys.argv
    destino = os.path.join(RAIZ, "demo", "matriz_demo.xlsx") if demo \
        else os.path.join(RAIZ, "plantilla", "Quorum_matriz.xlsx")
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    libro(demo).save(destino)
    print("Escrito:", os.path.relpath(destino, RAIZ))
